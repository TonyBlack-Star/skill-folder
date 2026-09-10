# -*- coding: utf-8 -*-
"""
把整合后的 XLSX 套入杂志风模板，生成可脱机打开的独立单文件看板。

用法：
    python make_dashboard.py
    python make_dashboard.py --xlsx r.xlsx --template t.html --out d.html --currency-symbol "$"

要点：
- 数据以 JSON 注入模板的 const DATA 占位符
- Chart.js 优先 CDN，失败时启用 --chartjs 指定的本地副本，保证脱机可用
- 额外生成「汇总」产品（全店日度聚合；计数求和、比率按合计重算）
- XLSX 比率列底层为小数（0-1），JS 端再 ×100 显示为 %
"""
import argparse
import json
import os
import re
from collections import defaultdict

from openpyxl import load_workbook

COUNTS = ("Online store visitors", "Sessions", "Sessions with cart additions",
          "Sessions that reached checkout", "Sessions that completed checkout",
          "Net items sold", "Total sales")
AGG_NAME = "汇总"


def build_data(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    products, daily, date_set = {}, [], set()

    for ws in wb.worksheets:
        handle = ws.title
        pid = (ws["B1"].value or "").replace("Product ID:", "").strip()
        products[handle] = pid
        header = [c.value for c in ws[2]]
        idx = {name: i for i, name in enumerate(header) if name}

        for r in range(3, ws.max_row + 1):
            day = ws.cell(r, idx["Day"] + 1).value
            if day is None:
                continue
            day = str(day)[:10]
            date_set.add(day)

            def cell(name):
                if name not in idx:
                    return None
                v = ws.cell(r, idx[name] + 1).value
                if v is None or v == "":
                    return None
                try:
                    return float(v)
                except (TypeError, ValueError):
                    return None

            sessions = cell("Sessions") or 0.0
            reached = cell("Sessions that reached checkout") or 0.0
            sold = cell("Net items sold") or 0.0
            daily.append({
                "Day": day, "Product": handle, "Product ID": pid,
                "Landing page path": "/products/" + handle,
                "Online store visitors": cell("Online store visitors"),
                "Sessions": cell("Sessions"),
                "Sessions with cart additions": cell("Sessions with cart additions"),
                "Sessions that reached checkout": cell("Sessions that reached checkout"),
                "Sessions that completed checkout": cell("Sessions that completed checkout"),
                "Bounce rate": (cell("Bounce rate") if sessions else None),
                "Added to cart rate": (cell("Added to cart rate") if sessions else None),
                "Reached checkout rate": (cell("Reached checkout rate") if sessions else None),
                "Checkout conversion rate": (cell("Checkout conversion rate") if reached else None),
                "Conversion rate": (cell("Conversion rate") if sessions else None),
                "Average session duration": cell("Average session duration"),
                "Net items sold": cell("Net items sold"),
                "Total sales": cell("Total sales"),
                "自定义转化率": (sold / sessions if sessions else None),
            })

    # 全店日度汇总：计数求和；比率由合计分子/分母重算；跳出率与时长加权平均
    agg = defaultdict(lambda: {**dict.fromkeys(COUNTS, 0.0),
                               "_bw": 0.0, "_bs": 0.0, "_dw": 0.0, "_ds": 0.0})
    for row in daily:
        a = agg[row["Day"]]
        s = row["Sessions"] or 0.0
        for k in COUNTS:
            if row.get(k) is not None:
                a[k] += row[k]
        if row.get("Bounce rate") is not None and s:
            a["_bw"] += row["Bounce rate"] * s
            a["_bs"] += s
        if row.get("Average session duration") is not None and s:
            a["_dw"] += row["Average session duration"] * s
            a["_ds"] += s

    for day in sorted(agg.keys()):
        a = agg[day]
        s = a["Sessions"]
        reached = a["Sessions that reached checkout"]
        completed = a["Sessions that completed checkout"]
        daily.append({
            "Day": day, "Product": AGG_NAME, "Product ID": "ALL_PRODUCTS",
            "Landing page path": "/products/all",
            "Online store visitors": a["Online store visitors"] or None,
            "Sessions": s or None,
            "Sessions with cart additions": a["Sessions with cart additions"] or None,
            "Sessions that reached checkout": reached or None,
            "Sessions that completed checkout": completed or None,
            "Bounce rate": (a["_bw"] / a["_bs"] if a["_bs"] else None),
            "Added to cart rate": (a["Sessions with cart additions"] / s if s else None),
            "Reached checkout rate": (reached / s if s else None),
            "Checkout conversion rate": (completed / reached if reached else None),
            "Conversion rate": (completed / s if s else None),
            "Average session duration": (a["_dw"] / a["_ds"] if a["_ds"] else None),
            "Net items sold": a["Net items sold"] or None,
            "Total sales": a["Total sales"] or None,
            "自定义转化率": (a["Net items sold"] / s if s else None),
        })

    products = {AGG_NAME: "ALL_PRODUCTS", **products}
    daily.sort(key=lambda x: (x["Product"], x["Day"]))
    dates = sorted(date_set)
    return {"products": products, "daily": daily, "date_min": dates[0], "date_max": dates[-1]}


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=os.getcwd())
    ap.add_argument("--xlsx", default=None)
    ap.add_argument("--template", default=os.path.join(here, "..", "assets", "dashboard_template.html"))
    ap.add_argument("--chartjs", default=None, help="Chart.js UMD 本地副本（脱机兜底）")
    ap.add_argument("--out", default=None)
    ap.add_argument("--currency-symbol", default="€")
    ap.add_argument("--currency-code", default="EUR")
    args = ap.parse_args()

    D = os.path.abspath(args.dir)
    xlsx = args.xlsx or os.path.join(D, "integrated_result.xlsx")
    out = args.out or os.path.join(D, "traffic-sales-dashboard.html")
    chartjs = args.chartjs or os.path.join(D, "chartjs.min.js")

    DATA = build_data(xlsx)
    tpl = open(args.template, encoding="utf-8").read()
    if "/*__DATA__*/" not in tpl or "/*__CHARTJS__*/" not in tpl:
        raise SystemExit("模板缺少 /*__DATA__*/ 或 /*__CHARTJS__*/ 占位符")

    tpl = re.sub(r'const CURRENCY_SYMBOL="[^"]*",\s*CURRENCY_CODE="[^"]*";',
                 'const CURRENCY_SYMBOL="%s", CURRENCY_CODE="%s";'
                 % (args.currency_symbol, args.currency_code), tpl, count=1)

    html = tpl.replace("/*__DATA__*/", json.dumps(DATA, ensure_ascii=False))
    if os.path.exists(chartjs):
        lib = open(chartjs, encoding="utf-8").read().replace("</script", "<\\/script")
        html = html.replace("/*__CHARTJS__*/",
                            'if(typeof Chart==="undefined"){\n' + lib
                            + "\nwindow.__chartFromInline=true;\n}")
        offline = True
    else:
        html = html.replace("/*__CHARTJS__*/", "")
        offline = False

    with open(out, "w", encoding="utf-8") as f:
        f.write(html)

    print("产品数:", len(DATA["products"]))
    print("daily 行数:", len(DATA["daily"]))
    print("日期范围:", DATA["date_min"], "~", DATA["date_max"])
    print("图表库脱机兜底:", "已内嵌" if offline else "无（仅 CDN）")
    print("输出:", out)


if __name__ == "__main__":
    main()
